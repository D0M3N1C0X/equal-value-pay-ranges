"""
Builds deliverables/equal_value_pay_ranges.xlsx: the pay-range model a reward team would use.

Every number is a live formula over the input sheets: change a factor weight, a spread, a
positioning or an Eurostat figure and the grades, the ranges, every employee's position and the
job-advertisement ranges recalculate. The Reconciliation sheet holds the pandas results and checks
each one against its formula.

Conventions as in pay-transparency-readiness-kit: blue text = input, black = formula, green =
link to another sheet, yellow fill = key assumption, Arial throughout, no dynamic arrays.
"""
import math
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.hyperlink import Hyperlink

import config as C
import model
from deterministic import normalise

OUTPUT = C.DELIVERABLES / "equal_value_pay_ranges.xlsx"
REPO = "https://github.com/D0M3N1C0X/equal-value-pay-ranges"

FONT = "Arial"
BLUE, GREEN, INK, MUTED, WHITE = "0000FF", "008000", "1B2430", "5F6B7A", "FFFFFF"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
YELLOW = PatternFill("solid", fgColor="FFFF00")
PCT, PCT1, COUNT, RATIO, NUM = "0.00%", "0.0%", "#,##0", "0.000", "0.0000"
EUR = '"€"#,##0;-"€"#,##0;"-"'
LOCAL = '#,##0;-#,##0;"-"'


def font(color=INK, bold=False, italic=False, size=10):
    return Font(name=FONT, color=color, bold=bold, italic=italic, size=size)


def put(ws, ref, value, *, color=INK, bold=False, italic=False, size=10, fmt=None, fill=None, wrap=False):
    cell = ws[ref]
    cell.value = value
    cell.font = font(color, bold, italic, size)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def header(ws, row, labels, start=1, height=30):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=label)
        c.font = font(WHITE, bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def title(ws, text, subtitle=None):
    put(ws, "A1", text, bold=True, size=14)
    if subtitle:
        put(ws, "A2", subtitle, color=MUTED, italic=True)


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def plain(x):
    if x is None:
        return None
    if hasattr(x, "item"):
        x = x.item()
    if isinstance(x, float) and math.isnan(x):
        return None
    return x


def rng(sheet, letter, first, last):
    return f"'{sheet}'!${letter}${first}:${letter}${last}"


class Model:
    def __init__(self, out: dict):
        self.o = out
        self.wb = Workbook()
        self.refs = {}
        self.checks = []
        self.countries = list(C.ENTITIES)
        self.grades = [g for _, g in C.GRADES]

    def check(self, area, item, value, ref):
        self.checks.append((area, item, plain(value), ref))

    def build(self, path: Path):
        wb = self.wb
        names = ["Cover", "Summary", "Structure", "Job ads", "Employees", "Roles", "Current bands", "Market",
                 "Settings", "Reconciliation"]
        wb.active.title = names[0]
        for n in names[1:]:
            wb.create_sheet(n)
        self.settings(wb["Settings"])
        self.market(wb["Market"])
        self.roles(wb["Roles"])
        self.current_bands(wb["Current bands"])
        self.structure_index(wb["Structure"])
        self.employees(wb["Employees"])
        self.structure_ranges(wb["Structure"])
        self.summary(wb["Summary"])
        self.job_ads(wb["Job ads"])
        self.reconciliation(wb["Reconciliation"])
        self.cover(wb["Cover"], names)
        for ws in wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.calculation.fullCalcOnLoad = True
        stamp = datetime(2026, 1, 1)
        wb.properties.creator = "Domenico Perroni"
        wb.properties.title = "Equal-Value Pay Ranges"
        wb.properties.created = wb.properties.modified = stamp
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        normalise(path)

    # ---- Settings ------------------------------------------------------------------------------
    def settings(self, ws):
        title(ws, "Settings", "Every assumption. Blue cells are inputs; yellow ones move the results most.")
        widths(ws, {"A": 3, "B": 34, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14, "I": 70})
        R = self.refs
        put(ws, "B4", "Job evaluation: factor weights (Article 4(4) criteria)", bold=True, size=11)
        header(ws, 5, ["Factor", "Weight"], start=2, height=18)
        for i, f in enumerate(C.FACTORS, start=6):
            put(ws, f"B{i}", f.replace("_", " ").capitalize())
            put(ws, f"C{i}", C.FACTOR_WEIGHTS[f], color=BLUE)
            R[f"w_{f}"] = f"Settings!$C${i}"
        put(ws, "I6", "From pay-transparency-readiness-kit: whole-number weights summing to 100, so points are exact.",
            color=MUTED, italic=True)

        put(ws, "B11", "Grades: categories of equal value, and the range policy", bold=True, size=11)
        header(ws, 12, ["Grade", "Lowest points", "Range spread", "Sets the country level (1 = yes)"], start=2, height=30)
        first = 13
        for i, (floor, g) in enumerate(C.GRADES, start=first):
            put(ws, f"B{i}", g, color=BLUE)
            put(ws, f"C{i}", floor, color=BLUE)
            put(ws, f"D{i}", C.SPREAD[g], color=BLUE, fmt=PCT1, fill=YELLOW)
            put(ws, f"E{i}", int(g in C.LEVEL_GRADES), color=BLUE)
        last = first + len(C.GRADES) - 1
        R["grades"], R["floors"] = f"Settings!$B${first}:$B${last}", f"Settings!$C${first}:$C${last}"
        R["spreads"], R["level_flags"] = f"Settings!$D${first}:$D${last}", f"Settings!$E${first}:$E${last}"
        put(ws, f"I{first}", "Spread = maximum / minimum - 1. Grades A-E set the country level: ISCO group 1 "
                              "averages every manager, so it is a poor yardstick for F and G.", color=MUTED, italic=True)
        r = last + 2
        put(ws, f"B{r}", "Reference grade (index = 1)")
        put(ws, f"C{r}", C.REFERENCE_GRADE, color=BLUE)
        R["ref_grade"] = f"Settings!$C${r}"
        r += 1
        put(ws, f"B{r}", "Minimum progression between midpoints")
        put(ws, f"C{r}", C.MIN_PROGRESSION, color=BLUE, fmt=PCT1, fill=YELLOW)
        put(ws, f"I{r}", "Grades cut across job levels, so the current bands can leave two adjacent grades almost "
                         "level; each midpoint is kept at least this far above the one below.", color=MUTED, italic=True)
        R["min_prog"] = f"Settings!$C${r}"

        r += 2
        put(ws, f"B{r}", "Market and currency", bold=True, size=11)
        r += 1
        put(ws, f"B{r}", "EUR/PLN")
        put(ws, f"C{r}", C.EUR_PLN, color=BLUE, fmt="0.0000")
        put(ws, f"I{r}", "ECB euro reference rate, monthly average June 2026.", color=MUTED, italic=True)
        R["eur_pln"] = f"Settings!$C${r}"
        r += 1
        put(ws, f"B{r}", "Update from 2025 to mid-2026, years")
        put(ws, f"C{r}", C.UPDATE_YEARS, color=BLUE, fmt="0.0", fill=YELLOW)
        put(ws, f"I{r}", "Half a year at each country's 2024-2025 labour cost growth. Illustrative.", color=MUTED,
            italic=True)
        R["update"] = f"Settings!$C${r}"
        r += 1
        put(ws, f"B{r}", "Monthly ad rounding step")
        put(ws, f"C{r}", C.AD_MONTHLY_STEP, color=BLUE, fmt=COUNT)
        R["month_step"] = f"Settings!$C${r}"

        r += 2
        put(ws, f"B{r}", "Countries: pay policy and job advertisements", bold=True, size=11)
        r += 1
        header(ws, r, ["Country", "Entity", "Currency", "Currency per euro", "Positioning vs market",
                       "Ad shows pay per", "Annual ad rounding"], start=2, height=30)
        r += 1
        first = r
        for c in self.countries:
            cur = C.CURRENCY[c]
            put(ws, f"B{r}", c, color=BLUE)
            put(ws, f"C{r}", C.ENTITIES[c], color=BLUE)
            put(ws, f"D{r}", cur, color=BLUE)
            put(ws, f"E{r}", f"={R['eur_pln'].split('!')[1]}" if cur == "PLN" else 1, color=INK if cur == "PLN" else BLUE,
                fmt="0.0000")
            put(ws, f"F{r}", C.POSITIONING[c], color=BLUE, fmt=PCT1, fill=YELLOW)
            put(ws, f"G{r}", C.AD_PERIOD[c], color=BLUE)
            put(ws, f"H{r}", C.AD_STEP[cur], color=BLUE, fmt=COUNT)
            r += 1
        last = r - 1
        for key, letter in (("c_codes", "B"), ("c_cur", "D"), ("c_fx", "E"), ("c_pos", "F"), ("c_period", "G"),
                            ("c_step", "H")):
            R[key] = f"Settings!${letter}${first}:${letter}${last}"
        put(ws, f"I{first}", "Positioning 100% = the midpoint pays what large employers pay on average for the same "
                              "occupations. Poland advertises monthly gross pay, the others annual.", color=MUTED,
            italic=True)

    def lookup(self, table, key_range, key):
        return f"INDEX({table},MATCH({key},{key_range},0))"

    # ---- Market ------------------------------------------------------------------------------------
    def market(self, ws):
        o, R = self.o, self.refs
        title(ws, "Market: Eurostat Structure of Earnings Survey 2022",
              "Mean annual gross earnings by occupation (ISCO-08 major group), enterprises with 10+ employees, "
              "NACE B-S excluding O, national currency. Blue figures are copied from the snapshots in data/raw/eurostat.")
        widths(ws, {"A": 16, "B": 10, "C": 10, **{col(k): 14 for k in range(4, 13)}})
        put(ws, "A4", "Labour cost index, wages and salaries (lc_lci_r2_a, 2020 = 100)", bold=True, size=11)
        header(ws, 5, ["Country", f"LCI {C.MARKET_YEAR}", f"LCI {C.INDEX_LATEST - 1}", f"LCI {C.INDEX_LATEST}",
                       "Growth, latest year", "Ageing factor"], height=30)
        lci = o["lci"]
        first = 6
        for i, c in enumerate(self.countries, start=first):
            put(ws, f"A{i}", c, color=BLUE)
            put(ws, f"B{i}", lci.loc[c, "lci_base"], color=BLUE, fmt="0.0")
            put(ws, f"C{i}", lci.loc[c, "lci_prev"], color=BLUE, fmt="0.0")
            put(ws, f"D{i}", lci.loc[c, "lci_latest"], color=BLUE, fmt="0.0")
            put(ws, f"E{i}", f"=D{i}/C{i}-1", fmt=PCT)
            put(ws, f"F{i}", f"=D{i}/B{i}*(D{i}/C{i})^{R['update']}", fmt=NUM)
            self.check("Market", f"{c} ageing factor", lci.loc[c, "aging"], f"Market!F{i}")
            self.check("Market", f"{c} labour cost growth", lci.loc[c, "growth"], f"Market!E{i}")
        last = first + len(self.countries) - 1
        R["lci_c"], R["lci_growth"], R["lci_aging"] = (f"Market!$A${first}:$A${last}", f"Market!$E${first}:$E${last}",
                                                       f"Market!$F${first}:$F${last}")

        r = last + 3
        put(ws, f"A{r}", "Large-employer premium: enterprises of 1,000+ (earn_ses22_32) against 10+ (earn_ses22_28)",
            bold=True, size=11)
        r += 1
        header(ws, r, ["Country", "Occupation", "Earnings, 10+", "Earnings, 1,000+", "Ratio"], height=30)
        r += 1
        first = r
        for t in o["premium_table"].itertuples(index=False):
            put(ws, f"A{r}", t.country, color=BLUE)
            put(ws, f"B{r}", t.isco08, color=BLUE)
            put(ws, f"C{r}", t.ern_10, color=BLUE, fmt=LOCAL)
            put(ws, f"D{r}", t.ern_1000, color=BLUE, fmt=LOCAL)
            put(ws, f"E{r}", f"=D{r}/C{r}", fmt=NUM)
            r += 1
        last = r - 1
        put(ws, f"A{r}", "Premium (median)", bold=True)
        put(ws, f"E{r}", f"=MEDIAN(E{first}:E{last})", fmt=NUM, bold=True)
        put(ws, f"F{r}", "Italy's 1,000+ figures are confidential, so one premium applies to all four.", color=MUTED,
            italic=True)
        R["premium"] = f"Market!$E${r}"
        self.check("Market", "Large-employer premium", o["premium"], f"Market!E{r}")

        r += 3
        put(ws, f"A{r}", "Market base pay by occupation, full time, per year", bold=True, size=11)
        r += 1
        header(ws, r, ["Key", "Country", "Occupation", "Earnings", "Annual bonuses", "Base pay", "Ageing",
                       "Premium", "Mid-2026, local", "Currency per euro", "Mid-2026, EUR"], height=30)
        r += 1
        first = r
        for t in o["market"].itertuples(index=False):
            put(ws, f"A{r}", f'=B{r}&"|"&C{r}')
            put(ws, f"B{r}", t.country, color=BLUE)
            put(ws, f"C{r}", t.isco08, color=BLUE)
            put(ws, f"D{r}", t.ern, color=BLUE, fmt=LOCAL)
            put(ws, f"E{r}", t.bns, color=BLUE, fmt=LOCAL)
            put(ws, f"F{r}", f"=D{r}-E{r}", fmt=LOCAL)
            put(ws, f"G{r}", "=" + self.lookup(R["lci_aging"], R["lci_c"], f"B{r}"), color=GREEN, fmt=NUM)
            put(ws, f"H{r}", f"={R['premium']}", color=GREEN, fmt=NUM)
            put(ws, f"I{r}", f"=F{r}*G{r}*H{r}", fmt=LOCAL)
            put(ws, f"J{r}", "=" + self.lookup(R["c_fx"], R["c_codes"], f"B{r}"), color=GREEN, fmt="0.0000")
            put(ws, f"K{r}", f"=I{r}/J{r}", fmt=EUR)
            self.check("Market", f"{t.country} {t.isco08} market pay, EUR", t.eur, f"Market!K{r}")
            r += 1
        last = r - 1
        R["mk_key"], R["mk_eur"] = f"Market!$A${first}:$A${last}", f"Market!$K${first}:$K${last}"
        ws.freeze_panes = "A6"

    # ---- Roles -------------------------------------------------------------------------------------
    def roles(self, ws):
        R, r_ = self.refs, self.o["roles"]
        title(ws, "Roles: job evaluation, grade and market occupation",
              "Scores 1-5 on the four Article 4(4) criteria (from pay-transparency-readiness-kit); points, grade and "
              "the market reference follow. The occupation mapping is a judgement: see data/isco_map.csv.")
        widths(ws, {"A": 22, "B": 22, "C": 18, "D": 8, **{col(k): 11 for k in range(5, 17)}})
        header(ws, 4, ["Key", "Role", "Department", "Level", "Skills", "Effort", "Responsibility", "Working conditions",
                       "Points", "Grade", "Occupation", *[f"Market {c}, EUR" for c in self.countries]], height=32)
        first = 5
        for i, t in enumerate(r_.itertuples(index=False), start=first):
            put(ws, f"A{i}", f'=C{i}&"|"&D{i}')
            put(ws, f"B{i}", t.role, color=BLUE)
            put(ws, f"C{i}", t.department, color=BLUE)
            put(ws, f"D{i}", t.job_level, color=BLUE)
            for j, f in enumerate(C.FACTORS):
                put(ws, f"{col(5 + j)}{i}", getattr(t, f), color=BLUE)
            put(ws, f"I{i}", f"=E{i}*{R['w_skills']}+F{i}*{R['w_effort']}+G{i}*{R['w_responsibility']}"
                             f"+H{i}*{R['w_working_conditions']}")
            put(ws, f"J{i}", f"=INDEX({R['grades']},MATCH(I{i},{R['floors']},1))")
            put(ws, f"K{i}", t.isco08, color=BLUE)
            mk = self.o["market"].set_index(["country", "isco08"])["eur"]
            for j, c in enumerate(self.countries):
                put(ws, f"{col(12 + j)}{i}", "=" + self.lookup(R["mk_eur"], R["mk_key"], f'"{c}|"&K{i}'),
                    color=GREEN, fmt=EUR)
                if i % 3 == 0:
                    self.check("Roles", f"{t.role} market {c}", mk[(c, t.isco08)], f"Roles!{col(12 + j)}{i}")
            self.check("Roles", f"{t.role} points", t.points, f"Roles!I{i}")
            self.check("Roles", f"{t.role} grade", t.grade, f"Roles!J{i}")
        last = first + len(r_) - 1
        R["role_key"], R["role_grade"], R["role_isco"] = (f"Roles!$A${first}:$A${last}", f"Roles!$J${first}:$J${last}",
                                                          f"Roles!$K${first}:$K${last}")
        ws.freeze_panes = "B5"

    # ---- Current bands -----------------------------------------------------------------------------
    def current_bands(self, ws):
        R, b = self.refs, self.o["bands"]
        title(ws, "Current bands: the company's midpoints by country, family and level (HRIS)",
              "One midpoint per role and country. The log feeds the grade index: family premiums average out.")
        widths(ws, {"A": 10, "B": 20, "C": 8, "D": 8, "E": 14, "F": 12})
        header(ws, 4, ["Country", "Department", "Level", "Grade", "Band midpoint", "Log"], height=20)
        first = 5
        for i, t in enumerate(b.itertuples(index=False), start=first):
            put(ws, f"A{i}", t.country, color=BLUE)
            put(ws, f"B{i}", t.department, color=BLUE)
            put(ws, f"C{i}", t.job_level, color=BLUE)
            put(ws, f"D{i}", "=" + self.lookup(R["role_grade"], R["role_key"], f'B{i}&"|"&C{i}'), color=GREEN)
            put(ws, f"E{i}", t.band_mid, color=BLUE, fmt=EUR)
            put(ws, f"F{i}", f"=LN(E{i})", fmt=NUM)
            if (i - first) % 6 == 0:
                self.check("Current bands", f"{t.country} {t.department} {t.job_level} grade", t.grade,
                           f"'Current bands'!D{i}")
        last = first + len(b) - 1
        R["cb_c"], R["cb_g"], R["cb_ln"] = (f"'Current bands'!$A${first}:$A${last}",
                                            f"'Current bands'!$D${first}:$D${last}",
                                            f"'Current bands'!$F${first}:$F${last}")
        ws.freeze_panes = "A5"

    # ---- Structure ---------------------------------------------------------------------------------
    def structure_index(self, ws):
        R, gi = self.refs, self.o["grade_index"]
        title(ws, "Structure: one range per grade and country",
              "Grade index from the current bands with family premiums averaged out; country level from the market; "
              "midpoint = level x index; minimum and maximum from the spread.")
        widths(ws, {"A": 12, "B": 10, **{col(k): 13 for k in range(3, 15)}})
        put(ws, "A4", "Grade index: current progression, family-neutral", bold=True, size=11)
        n = len(self.countries)
        header(ws, 5, ["Grade", *[f"Log mean {c}" for c in self.countries], *[f"vs {C.REFERENCE_GRADE}, {c}" for c in self.countries],
                       "Index, current", "Index"], height=30)
        first = 6
        last = first + len(self.grades) - 1
        ref_row = first + self.grades.index(C.REFERENCE_GRADE)
        for i, g in enumerate(self.grades, start=first):
            put(ws, f"A{i}", g, color=BLUE)
            for j, c in enumerate(self.countries):
                put(ws, f"{col(2 + j)}{i}", f'=AVERAGEIFS({R["cb_ln"]},{R["cb_c"]},"{c}",{R["cb_g"]},A{i})', fmt=NUM)
                put(ws, f"{col(2 + n + j)}{i}", f"={col(2 + j)}{i}-{col(2 + j)}${ref_row}", fmt=NUM)
            raw, sm = col(2 + 2 * n), col(3 + 2 * n)
            put(ws, f"{raw}{i}", f"=EXP(AVERAGE({col(2 + n)}{i}:{col(1 + 2 * n)}{i}))", fmt=RATIO)
            put(ws, f"{sm}{i}", f"={raw}{i}" if i == first else f"=MAX({raw}{i},{sm}{i - 1}*(1+{R['min_prog']}))",
                fmt=RATIO, bold=True)
            self.check("Structure", f"Grade {g} index, current", self.o["raw_index"][g], f"Structure!{raw}{i}")
            self.check("Structure", f"Grade {g} index", gi[g], f"Structure!{sm}{i}")
        sm = col(3 + 2 * n)
        R["gi_g"], R["gi"] = f"Structure!$A${first}:$A${last}", f"Structure!${sm}${first}:${sm}${last}"
        self.row = last + 3
        # rows of the ranges table written by structure_ranges(), known now so the Employees sheet can use them
        st_first = self.row + 10
        st_last = st_first + len(self.countries) * len(self.grades) - 1
        for key, letter in (("st_key", "A"), ("st_min", "G"), ("st_mid", "H"), ("st_max", "I"),
                            ("st_min_l", "K"), ("st_max_l", "M")):
            R[key] = f"Structure!${letter}${st_first}:${letter}${st_last}"
        self.st_first = st_first

    def employees(self, ws):
        R, p = self.refs, self.o["positions"]
        title(ws, "Employees on the payroll at 30 June 2026 against the new ranges",
              "Inputs in blue from the HRIS; everything else is a formula. Salaries are full-time annual base pay in euro.")
        widths(ws, {"A": 10, "B": 8, "C": 18, "D": 7, "E": 7, "F": 6, **{col(k): 11 for k in range(7, 24)}})
        header(ws, 4, ["Employee", "Country", "Department", "Level", "Gender", "FTE", "Salary", "Current band mid",
                       "Grade", "Occupation", "Market", "Sets level", "Grade index", "New min", "New mid", "New max",
                       "Compa, current", "Compa, new", "Range penetration", "Status", "Cost to minimum",
                       "Above maximum", "Years to absorb"], height=32)
        first = 5
        last = first + len(p) - 1
        self.emp_rows = (first, last)
        for i, t in enumerate(p.itertuples(index=False), start=first):
            put(ws, f"A{i}", t.employee_id, color=BLUE)
            put(ws, f"B{i}", t.country, color=BLUE)
            put(ws, f"C{i}", t.department, color=BLUE)
            put(ws, f"D{i}", t.job_level, color=BLUE)
            put(ws, f"E{i}", t.gender, color=BLUE)
            put(ws, f"F{i}", t.fte, color=BLUE)
            put(ws, f"G{i}", t.salary, color=BLUE, fmt=EUR)
            put(ws, f"H{i}", t.band_mid, color=BLUE, fmt=EUR)
            put(ws, f"I{i}", "=" + self.lookup(R["role_grade"], R["role_key"], f'C{i}&"|"&D{i}'), color=GREEN)
            put(ws, f"J{i}", "=" + self.lookup(R["role_isco"], R["role_key"], f'C{i}&"|"&D{i}'), color=GREEN)
            put(ws, f"K{i}", "=" + self.lookup(R["mk_eur"], R["mk_key"], f'B{i}&"|"&J{i}'), color=GREEN, fmt=EUR)
            put(ws, f"L{i}", "=" + self.lookup(R["level_flags"], R["grades"], f"I{i}"), color=GREEN)
            put(ws, f"M{i}", "=" + self.lookup(R["gi"], R["gi_g"], f"I{i}"), color=GREEN, fmt=RATIO)
            for letter, key in (("N", "st_min"), ("O", "st_mid"), ("P", "st_max")):
                put(ws, f"{letter}{i}", "=" + self.lookup(R[key], R["st_key"], f'B{i}&"|"&I{i}'), color=GREEN, fmt=EUR)
            put(ws, f"Q{i}", f"=G{i}/H{i}", fmt=RATIO)
            put(ws, f"R{i}", f"=G{i}/O{i}", fmt=RATIO)
            put(ws, f"S{i}", f"=(G{i}-N{i})/(P{i}-N{i})", fmt=PCT1)
            put(ws, f"T{i}", f'=IF(G{i}<N{i},"Below",IF(G{i}>P{i},"Above","Within"))')
            put(ws, f"U{i}", f"=MAX(0,N{i}-G{i})*F{i}", fmt=EUR)
            put(ws, f"V{i}", f"=MAX(0,G{i}-P{i})*F{i}", fmt=EUR)
            put(ws, f"W{i}", f'=IF(T{i}="Above",LN(G{i}/P{i})/LN(1+{self.lookup(R["lci_growth"], R["lci_c"], f"B{i}")}),0)',
                fmt="0.0")
        step = max(1, len(p) // 40)
        for k in list(range(0, len(p), step))[:40]:
            i, t = first + k, p.iloc[k]
            for letter, name, v in (("R", "compa new", t.compa_new), ("T", "status", t.status),
                                    ("U", "cost to minimum", t.to_min), ("V", "above maximum", t.above_max),
                                    ("W", "years to absorb", t.years_to_absorb)):
                self.check("Employees", f"{t.employee_id} {name}", v, f"Employees!{letter}{i}")
        E = lambda letter: f"Employees!${letter}${first}:${letter}${last}"
        R.update({"e_c": E("B"), "e_dep": E("C"), "e_gen": E("E"), "e_sal": E("G"), "e_mkt": E("K"), "e_flag": E("L"),
                  "e_idx": E("M"), "e_cc": E("Q"), "e_cn": E("R"), "e_st": E("T"), "e_min": E("U"), "e_max": E("V")})
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:W{last}"

    def structure_ranges(self, ws):
        R, o = self.refs, self.o
        r = self.row
        put(ws, f"A{r}", "Country level: the market for the non-managerial workforce", bold=True, size=11)
        r += 1
        header(ws, r, ["Country", "Market pay, grades A-E", "Grade index, grades A-E", "Positioning", "Level (grade D mid)"],
               height=30)
        r += 1
        first = r
        st = o["structure"].set_index(["country", "grade"])
        for c in self.countries:
            put(ws, f"A{r}", c, color=BLUE)
            put(ws, f"B{r}", f'=SUMIFS({R["e_mkt"]},{R["e_c"]},A{r},{R["e_flag"]},1)', fmt=EUR)
            put(ws, f"C{r}", f'=SUMIFS({R["e_idx"]},{R["e_c"]},A{r},{R["e_flag"]},1)', fmt="#,##0.00")
            put(ws, f"D{r}", "=" + self.lookup(R["c_pos"], R["c_codes"], f"A{r}"), color=GREEN, fmt=PCT1)
            put(ws, f"E{r}", f"=D{r}*B{r}/C{r}", fmt=EUR, bold=True)
            self.check("Structure", f"{c} level", st.loc[(c, C.REFERENCE_GRADE), "level"], f"Structure!E{r}")
            r += 1
        R["lv_c"], R["lv"] = f"Structure!$A${first}:$A${r - 1}", f"Structure!$E${first}:$E${r - 1}"

        r += 2
        put(ws, f"A{r}", "Ranges", bold=True, size=11)
        r += 1
        header(ws, r, ["Key", "Country", "Grade", "Index", "Level", "Spread", "Minimum", "Midpoint", "Maximum",
                       "Currency per euro", "Minimum, local", "Midpoint, local", "Maximum, local"], height=30)
        r += 1
        first = r
        assert first == self.st_first
        for c in self.countries:
            for g in self.grades:
                put(ws, f"A{r}", f'=B{r}&"|"&C{r}')
                put(ws, f"B{r}", c, color=BLUE)
                put(ws, f"C{r}", g, color=BLUE)
                put(ws, f"D{r}", "=" + self.lookup(R["gi"], R["gi_g"], f"C{r}"), color=GREEN, fmt=RATIO)
                put(ws, f"E{r}", "=" + self.lookup(R["lv"], R["lv_c"], f"B{r}"), color=GREEN, fmt=EUR)
                put(ws, f"F{r}", "=" + self.lookup(R["spreads"], R["grades"], f"C{r}"), color=GREEN, fmt=PCT1)
                put(ws, f"H{r}", f"=E{r}*D{r}", fmt=EUR, bold=True)
                put(ws, f"G{r}", f"=H{r}*2/(2+F{r})", fmt=EUR)
                put(ws, f"I{r}", f"=G{r}*(1+F{r})", fmt=EUR)
                put(ws, f"J{r}", "=" + self.lookup(R["c_fx"], R["c_codes"], f"B{r}"), color=GREEN, fmt="0.0000")
                put(ws, f"K{r}", f"=G{r}*J{r}", fmt=LOCAL)
                put(ws, f"L{r}", f"=H{r}*J{r}", fmt=LOCAL)
                put(ws, f"M{r}", f"=I{r}*J{r}", fmt=LOCAL)
                for letter, key in (("G", "min"), ("H", "mid"), ("I", "max")):
                    self.check("Structure", f"{c} grade {g} {key}", st.loc[(c, g), key], f"Structure!{letter}{r}")
                self.check("Structure", f"{c} grade {g} max, local", st.loc[(c, g), "max_local"], f"Structure!M{r}")
                r += 1
        # range chart per country
        for k, c in enumerate(self.countries):
            chart = BarChart()
            chart.type, chart.grouping, chart.overlap = "bar", "stacked", 100
            chart.title = f"{C.ENTITIES[c]}: ranges by grade, EUR"
            chart.height, chart.width = 7, 11
            top = first + k * len(self.grades)
            helper_first = top
            for j, g in enumerate(self.grades):
                rr = top + j
                put(ws, f"O{rr}", g, color=MUTED)
                put(ws, f"P{rr}", f"=G{rr}", color=MUTED, fmt=EUR)
                put(ws, f"Q{rr}", f"=I{rr}-G{rr}", color=MUTED, fmt=EUR)
            data = Reference(ws, min_col=16, max_col=17, min_row=helper_first, max_row=helper_first + len(self.grades) - 1)
            chart.add_data(data, titles_from_data=False)
            chart.set_categories(Reference(ws, min_col=15, min_row=helper_first, max_row=helper_first + len(self.grades) - 1))
            chart.series[0].graphicalProperties.noFill = True
            chart.series[1].graphicalProperties.solidFill = "2A78D6"
            chart.legend = None
            ws.add_chart(chart, f"S{4 + k * 15}")
        put(ws, f"O{first - 1}", "Chart data: minimum and width", color=MUTED, italic=True)
        ws.freeze_panes = "A6"

    # ---- Summary -----------------------------------------------------------------------------------
    def summary(self, ws):
        R, s = self.refs, self.o["summary"].set_index("country")
        title(ws, "Summary: the workforce against the new ranges",
              "Payroll against market on grades A-E; positions and costs for everyone on the payroll at 30 June 2026.")
        widths(ws, {"A": 34, **{col(k): 15 for k in range(2, 8)}})
        header(ws, 4, ["Measure", *[C.ENTITIES[c] for c in self.countries], "All"], height=30)
        rows = [
            ("Headcount", "headcount", lambda c: f'=COUNTIF({R["e_c"]},"{c}")', COUNT, True),
            ("Payroll, grades A-E", "payroll", lambda c: f'=SUMIFS({R["e_sal"]},{R["e_c"]},"{c}",{R["e_flag"]},1)', EUR, True),
            ("Market for the same occupations", "market", lambda c: f'=SUMIFS({R["e_mkt"]},{R["e_c"]},"{c}",{R["e_flag"]},1)', EUR, True),
            ("Payroll against market", "ratio", None, RATIO, False),
            ("Average compa-ratio, current bands", "compa_current", lambda c: f'=AVERAGEIFS({R["e_cc"]},{R["e_c"]},"{c}")', RATIO, False),
            ("Average compa-ratio, new ranges", "compa_new", lambda c: f'=AVERAGEIFS({R["e_cn"]},{R["e_c"]},"{c}")', RATIO, False),
            ("Below the new minimum", "below", lambda c: f'=COUNTIFS({R["e_c"]},"{c}",{R["e_st"]},"Below")', COUNT, True),
            ("Within the range", "within", lambda c: f'=COUNTIFS({R["e_c"]},"{c}",{R["e_st"]},"Within")', COUNT, True),
            ("Above the new maximum", "above", lambda c: f'=COUNTIFS({R["e_c"]},"{c}",{R["e_st"]},"Above")', COUNT, True),
            ("Women", "women", lambda c: f'=COUNTIFS({R["e_c"]},"{c}",{R["e_gen"]},"F")', COUNT, True),
            ("Women below the minimum", "below_women", lambda c: f'=COUNTIFS({R["e_c"]},"{c}",{R["e_gen"]},"F",{R["e_st"]},"Below")', COUNT, True),
            ("Men", "men", lambda c: f'=COUNTIFS({R["e_c"]},"{c}",{R["e_gen"]},"M")', COUNT, True),
            ("Men below the minimum", "below_men", lambda c: f'=COUNTIFS({R["e_c"]},"{c}",{R["e_gen"]},"M",{R["e_st"]},"Below")', COUNT, True),
            ("Cost to bring everyone to the minimum, per year", "to_min", lambda c: f'=SUMIFS({R["e_min"]},{R["e_c"]},"{c}")', EUR, True),
            ("of which to women", "to_min_women", lambda c: f'=SUMIFS({R["e_min"]},{R["e_c"]},"{c}",{R["e_gen"]},"F")', EUR, True),
            ("Pay above the new maximum, per year", "above_max", lambda c: f'=SUMIFS({R["e_max"]},{R["e_c"]},"{c}")', EUR, True),
        ]
        r = 5
        where = {}
        for label, key, f, fmt, total in rows:
            put(ws, f"A{r}", label, bold=key in ("ratio", "to_min", "above_max"))
            where[key] = r
            for j, c in enumerate(self.countries):
                letter = col(2 + j)
                if key == "ratio":
                    formula = f"={letter}{where['payroll']}/{letter}{where['market']}"
                else:
                    formula = f(c)
                put(ws, f"{letter}{r}", formula, fmt=fmt)
                self.check("Summary", f"{c} {label.lower()}", s.loc[c, key], f"Summary!{letter}{r}")
            if total:
                put(ws, f"F{r}", f"=SUM(B{r}:E{r})", fmt=fmt, bold=True)
            elif key == "ratio":
                put(ws, f"F{r}", f"=F{where['payroll']}/F{where['market']}", fmt=fmt, bold=True)
            r += 1
        r += 2
        put(ws, f"A{r}", "Average compa-ratio against the new ranges, by family", bold=True, size=11)
        r += 1
        header(ws, r, ["Department", *[C.ENTITIES[c] for c in self.countries]], height=20)
        r += 1
        fam = self.o["family"]
        for d in C.DEPARTMENTS:
            put(ws, f"A{r}", d)
            for j, c in enumerate(self.countries):
                letter = col(2 + j)
                put(ws, f"{letter}{r}", f'=AVERAGEIFS({R["e_cn"]},{R["e_c"]},"{c}",{R["e_dep"]},A{r})', fmt=RATIO)
                self.check("Summary", f"{c} {d} compa-ratio", fam.loc[d, c], f"Summary!{letter}{r}")
            r += 1
        put(ws, f"A{r + 1}", "Above 1: the family is paid above the equal-value midpoint, a premium the new structure no "
                             "longer carries.", color=MUTED, italic=True)
        ws.freeze_panes = "B5"

    # ---- Job ads -----------------------------------------------------------------------------------
    def job_ads(self, ws):
        R, a = self.refs, self.o["ads"]
        title(ws, "Job advertisements: the pay range to publish (Article 5)",
              "The grade's range in local currency, rounded outwards; Poland advertises monthly gross pay, the others "
              "annual. Full-time equivalents.")
        widths(ws, {"A": 26, "B": 22, "C": 7, "D": 8, "E": 9, "F": 8, **{col(k): 12 for k in range(7, 13)}, "M": 52})
        header(ws, 4, ["Key", "Role", "Grade", "Country", "Currency", "Per", "Minimum, local", "Maximum, local",
                       "Annual from", "Annual to", "Monthly from", "Monthly to", "Text for the advertisement"], height=30)
        r = 5
        for t in a.itertuples(index=False):
            put(ws, f"A{r}", f'=D{r}&"|"&C{r}')
            put(ws, f"B{r}", t.role, color=BLUE)
            dep, lev = t.department, t.job_level
            put(ws, f"C{r}", "=" + self.lookup(R["role_grade"], R["role_key"], f'"{dep}|{lev}"'), color=GREEN)
            put(ws, f"D{r}", t.country, color=BLUE)
            put(ws, f"E{r}", "=" + self.lookup(R["c_cur"], R["c_codes"], f"D{r}"), color=GREEN)
            put(ws, f"F{r}", "=" + self.lookup(R["c_period"], R["c_codes"], f"D{r}"), color=GREEN)
            put(ws, f"G{r}", "=" + self.lookup(R["st_min_l"], R["st_key"], f"A{r}"), color=GREEN, fmt=LOCAL)
            put(ws, f"H{r}", "=" + self.lookup(R["st_max_l"], R["st_key"], f"A{r}"), color=GREEN, fmt=LOCAL)
            step = self.lookup(R["c_step"], R["c_codes"], f"D{r}")
            put(ws, f"I{r}", f"=ROUNDDOWN(G{r}/{step},0)*{step}", fmt=LOCAL)
            put(ws, f"J{r}", f"=ROUNDUP(H{r}/{step},0)*{step}", fmt=LOCAL)
            put(ws, f"K{r}", f"=ROUNDDOWN(G{r}/12/{R['month_step']},0)*{R['month_step']}", fmt=LOCAL)
            put(ws, f"L{r}", f"=ROUNDUP(H{r}/12/{R['month_step']},0)*{R['month_step']}", fmt=LOCAL)
            put(ws, f"M{r}", (f'=E{r}&" "&TEXT(IF(F{r}="month",K{r},I{r}),"#,##0")&" – "'
                              f'&TEXT(IF(F{r}="month",L{r},J{r}),"#,##0")&" gross per "&F{r}&", full time"'))
            for letter, key in (("I", "annual_min"), ("J", "annual_max"), ("K", "monthly_min"), ("L", "monthly_max")):
                self.check("Job ads", f"{t.role} {t.country} {key.replace('_', ' ')}", getattr(t, key), f"'Job ads'!{letter}{r}")
            self.check("Job ads", f"{t.role} {t.country} text", t.text, f"'Job ads'!M{r}")
            r += 1
        ws.freeze_panes = "C5"
        ws.auto_filter.ref = f"A4:M{r - 1}"

    # ---- Reconciliation and cover ------------------------------------------------------------------
    def reconciliation(self, ws):
        title(ws, "Reconciliation: workbook formulas against the Python pipeline",
              "Column D was written by src/build_workbook.py from the pandas results; column E is the live formula.")
        widths(ws, {"A": 16, "B": 58, "C": 3, "D": 22, "E": 22, "F": 14, "G": 8})
        header(ws, 6, ["Area", "Item", "", "Python value", "Workbook value", "Difference", "Match"])
        first, last = 7, 6 + len(self.checks)
        put(ws, "B3", (f'=IF(COUNTIF(G{first}:G{last},"No")=0,"All "&COUNTA(B{first}:B{last})&" checks match",'
                       f'COUNTIF(G{first}:G{last},"No")&" of "&COUNTA(B{first}:B{last})&" checks do not match")'),
            bold=True, size=12)
        for r, (area, item, value, ref) in enumerate(self.checks, start=first):
            put(ws, f"A{r}", area)
            put(ws, f"B{r}", item)
            fmt = "0.000000" if isinstance(value, float) else None
            put(ws, f"D{r}", value, color=BLUE, fmt=fmt)
            put(ws, f"E{r}", f"={ref}", color=GREEN, fmt=fmt)
            put(ws, f"F{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")', fmt="0.0E+00")
            put(ws, f"G{r}", (f'=IF(AND(D{r}="",E{r}=""),"Yes",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),'
                              f'IF(ABS(E{r}-D{r})<=1E-9*MAX(1,ABS(D{r})),"Yes","No"),IF(D{r}=E{r},"Yes","No")))'))
        for text, colour in (("Yes", "D5ECDC"), ("No", "F5C6C2")):
            ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                                         fill=PatternFill("solid", fgColor=colour)))
        ws.freeze_panes = "A7"
        self.recon_result = "Reconciliation!B3"

    def cover(self, ws, names):
        widths(ws, {"A": 3, "B": 24, "C": 100})
        put(ws, "B2", "Equal-Value Pay Ranges", bold=True, size=18)
        put(ws, "B3", "From family-based bands to one range per grade and country, priced on Eurostat market data",
            color=MUTED, size=12)
        put(ws, "B5", "What it does", bold=True, size=11)
        lines = [
            "Grades 36 roles on the four Article 4(4) criteria, so roles of equal value share a grade across families.",
            "Prices the market for each role's occupation in Italy, Poland, Germany and Spain from Eurostat's 2022 "
            "earnings survey, aged to mid-2026, with a large-employer premium.",
            "Builds one range per grade and country: the current progression between grades, the market level, a "
            "policy spread.",
            "Places every employee against the new range, costs the move, and prints the range for each job "
            "advertisement (Article 5).",
        ]
        for i, text in enumerate(lines, start=6):
            put(ws, f"C{i}", text)
        put(ws, "B11", "How to read it", bold=True, size=11)
        for i, (label, meaning, colour, fill) in enumerate([
            ("Blue text", "an input", BLUE, None), ("Black text", "a formula", INK, None),
            ("Green text", "a link to another sheet", GREEN, None),
            ("Yellow fill", "a key assumption: spreads, positioning, the update to mid-2026", INK, YELLOW)], start=12):
            put(ws, f"B{i}", label, color=colour, fill=fill)
            put(ws, f"C{i}", meaning)
        put(ws, "B17", "Sheets", bold=True, size=11)
        purpose = {
            "Summary": "Payroll against market, positions against the new ranges, costs, by country, gender and family.",
            "Structure": "Grade index, country level and the 28 ranges, with a chart per country.",
            "Job ads": "The range to publish for every role in every country, and the sentence for the advertisement.",
            "Employees": "One row per employee: current band, new range, compa-ratio, status, cost.",
            "Roles": "Job evaluation scores, points, grade and market occupation for the 36 roles.",
            "Current bands": "The company's midpoints today, by country, family and level.",
            "Market": "Eurostat earnings, labour cost index and large-employer premium.",
            "Settings": "Factor weights, grades, spreads, positioning, currency and advertisement conventions.",
            "Reconciliation": "Each figure checked against the Python pipeline.",
        }
        for i, n in enumerate(names[1:], start=18):
            cell = put(ws, f"B{i}", n, color="1F5FA8")
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{n}'!A1")
            put(ws, f"C{i}", purpose[n])
        r = 18 + len(names)
        put(ws, f"B{r}", "Check", bold=True, size=11)
        put(ws, f"C{r}", f"={self.recon_result}", color=GREEN, bold=True)
        put(ws, f"B{r + 2}", "Data", bold=True, size=11)
        put(ws, f"C{r + 2}", "Employees are synthetic: the organisation analysed in hr-people-analytics, with its current "
                             "bands. Market figures are real: Eurostat, retrieved on the dates in data/raw/eurostat.",
            wrap=True)
        ws.row_dimensions[r + 2].height = 28
        put(ws, f"B{r + 3}", "Limits", bold=True, size=11)
        put(ws, f"C{r + 3}", "Occupations are ISCO major groups and earnings are means, not medians: a real engagement "
                             "prices jobs on salary surveys matched by job and level. See docs/method.md.", wrap=True)
        ws.row_dimensions[r + 3].height = 28
        put(ws, f"B{r + 5}", "Domenico Perroni", color=MUTED)
        cell = put(ws, f"C{r + 5}", REPO, color="1F5FA8")
        cell.hyperlink = REPO


def build(out: dict | None = None, path: Path = OUTPUT) -> Model:
    out = model.run() if out is None else out
    m = Model(out)
    m.build(path)
    return m


def main() -> None:
    m = build()
    print(f"workbook: {len(m.checks)} reconciliation checks -> {OUTPUT.relative_to(OUTPUT.parents[1])}")


if __name__ == "__main__":
    main()
